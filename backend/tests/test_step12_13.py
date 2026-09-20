"""Regression and local-storage security coverage for Steps 12 and 13."""

from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from pydantic import ValidationError
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

from app import documents as documents_module
from app import processing as processing_module
from app.content_models import CodeBlockRecord, ParsedBlock, ParsedBlockCollection, WorkbenchResult
from app.main import app
from app.parser import ParsedDocument
from app.schemas import DocumentDetail
from app.settings import settings
from app.task_store import (
    create_processing_task,
    get_document_task,
    processing_tasks,
    recover_interrupted_processing_tasks,
    update_processing_task,
)
from app.ai import AIProcessingError, TranslationOutput
from app.translation import prepare_translation


PDF_BYTES = b"%PDF-1.7\n1 0 obj << /Type /Catalog >> endobj\n%%EOF\n"


def make_two_column_pdf() -> bytes:
    content = (
        b"BT /F1 12 Tf 50 740 Td (LEFT-COLUMN-ONE) Tj "
        b"0 -20 Td (LEFT-COLUMN-TWO) Tj ET\n"
        b"BT /F1 12 Tf 320 740 Td (RIGHT-COLUMN-ONE) Tj "
        b"0 -20 Td (RIGHT-COLUMN-TWO) Tj ET"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n"
        + content + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(value)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(output)


def make_scanned_pdf(image_bytes: bytes, width: int, height: int) -> bytes:
    content = f"q {width} 0 0 {height} 0 0 cm /Im1 Do Q".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
            "/Resources << /XObject << /Im1 6 0 R >> >> /Contents 5 0 R >>"
        ).encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n"
        + content + b"\nendstream",
        (
            f"<< /Type /XObject /Subtype /Image /Width {width} /Height {height} "
            "/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode /Length "
        ).encode("ascii")
        + str(len(image_bytes)).encode("ascii")
        + b" >>\nstream\n"
        + image_bytes
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(value)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(output)


def make_docx() -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        archive.writestr(
            "word/document.xml",
            "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\"><w:body/></w:document>",
        )
    return buffer.getvalue()


class Step12And13Tests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="course-assistant-tests-")
        self.addCleanup(self.temp_dir.cleanup)
        root = Path(self.temp_dir.name)
        settings.upload_dir = root / "uploads"
        settings.parsed_dir = root / "parsed"
        settings.max_upload_size_mb = 50
        documents_module.documents.clear()
        processing_tasks.clear()
        self.client = TestClient(app)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)

    def parsed_document(self, document_id: str, *, with_code: bool = False) -> ParsedDocument:
        blocks: list[ParsedBlock] = []
        codes: list[CodeBlockRecord] = []
        if with_code:
            block = ParsedBlock(
                id=f"{document_id}:block:1",
                document_id=document_id,
                order=0,
                kind="code",
                docling_label="Code",
                text="print('safe')",
                original_text="print('safe')",
                code_language="python",
            )
            blocks.append(block)
            codes.append(
                CodeBlockRecord(
                    id=f"code:{document_id}",
                    block_id=block.id,
                    code="print('safe')",
                    language="python",
                )
            )
        collection = ParsedBlockCollection(document_id=document_id, blocks=blocks)
        workbench = WorkbenchResult(
            document_id=document_id,
            parsed_blocks=blocks,
            code_blocks=codes,
        )
        return ParsedDocument(
            markdown="# Sample course material\n",
            structured_data={"name": "safe sample"},
            status="success",
            parsed_blocks=collection,
            workbench_result=workbench,
        )

    def upload(self, filename: str = "sample.pdf", content: bytes = PDF_BYTES):
        return self.client.post(
            "/api/documents",
            files={"file": (filename, content, "application/octet-stream")},
        )

    def test_pdf_docx_upload_and_identical_uploads_get_independent_ids(self) -> None:
        with patch.object(
            documents_module,
            "parse_document",
            side_effect=lambda _path, document_id: self.parsed_document(document_id),
        ):
            first = self.upload()
            second = self.upload()
            docx = self.upload("table.docx", make_docx())

        self.assertEqual(first.status_code, 202)
        self.assertEqual(first.json()["status"], "pending")
        self.assertEqual(first.json()["processing_stage"], "queued")
        self.assertEqual(second.status_code, 202)
        self.assertNotEqual(first.json()["id"], second.json()["id"])
        self.assertEqual(docx.status_code, 202)
        self.assertEqual(docx.json()["file_type"], "docx")
        self.assertEqual(
            self.client.get(f"/api/documents/{first.json()['id']}/status").json()["status"],
            "completed",
        )

    def test_rejects_wrong_type_bad_pdf_bad_docx_and_oversize(self) -> None:
        with patch.object(documents_module, "parse_document"):
            self.assertEqual(self.upload("program.exe", b"MZ").status_code, 400)
            self.assertEqual(self.upload("wrong.pdf", b"not a pdf").status_code, 400)
            self.assertEqual(self.upload("wrong.docx", b"not a zip").status_code, 400)
            settings.max_upload_size_mb = 1
            response = self.upload("large.pdf", b"%PDF-1.7\n" + b"x" * (1024 * 1024))
            self.assertEqual(response.status_code, 413)
            self.assertEqual(response.json()["error"]["code"], "FILE_TOO_LARGE")

            unsafe_docx = BytesIO()
            with ZipFile(unsafe_docx, "w") as archive:
                archive.writestr("[Content_Types].xml", "<Types/>")
                archive.writestr("word/document.xml", "<document/>")
                archive.writestr("../outside.xml", "not allowed")
            self.assertEqual(self.upload("unsafe.docx", unsafe_docx.getvalue()).status_code, 400)

    def test_parse_failure_is_saved_and_can_be_retried(self) -> None:
        with patch.object(documents_module, "parse_document", side_effect=RuntimeError("private input")):
            uploaded = self.upload()
        document_id = uploaded.json()["id"]
        failed = self.client.get(f"/api/documents/{document_id}/status").json()
        self.assertEqual(failed["status"], "failed")
        self.assertIn("RuntimeError", failed["error_message"])
        self.assertNotIn("private input", failed["error_message"])

        with patch.object(
            documents_module,
            "parse_document",
            side_effect=lambda _path, current_id: self.parsed_document(current_id),
        ):
            retried = self.client.post(f"/api/documents/{document_id}/regenerate")
        self.assertEqual(retried.status_code, 202)
        self.assertEqual(retried.json()["status"], "pending")
        self.assertEqual(
            self.client.get(f"/api/documents/{document_id}/status").json()["status"],
            "completed",
        )

    def test_task_status_persists_errors_supports_retry_and_is_document_scoped(self) -> None:
        secret = "key-that-must-never-be-written"
        with patch.object(
            documents_module,
            "parse_document",
            side_effect=lambda _path, document_id: self.parsed_document(document_id, with_code=True),
        ):
            uploaded = self.upload()
        document_id = uploaded.json()["id"]
        with patch.object(
            processing_module,
            "explain_code_blocks",
            side_effect=AIProcessingError("AI_PROVIDER_UNAVAILABLE", "服务暂时无法处理请求"),
        ):
            started = self.client.post(
                f"/api/documents/{document_id}/process",
                json={"task": "explain_code"},
                headers={"X-User-API-Key": secret},
            )
        self.assertEqual(started.status_code, 202)
        task_id = started.json()["id"]
        status = self.client.get(f"/api/documents/{document_id}/tasks/{task_id}")
        self.assertEqual(status.json()["status"], "failed")
        self.assertEqual(status.json()["error_code"], "AI_PROVIDER_UNAVAILABLE")

        other_document_id = "123e4567-e89b-42d3-a456-426614174000"
        self.assertEqual(
            self.client.get(f"/api/documents/{other_document_id}/tasks/{task_id}").status_code,
            404,
        )
        self.assertEqual(
            self.client.get(f"/api/documents/{document_id}/tasks/../../task.json").status_code,
            404,
        )

        with self.assertLogs(processing_module.__name__, level="ERROR") as log_capture:
            with patch.object(
                processing_module,
                "explain_code_blocks",
                new_callable=AsyncMock,
                side_effect=RuntimeError(secret),
            ):
                unexpected = self.client.post(
                    f"/api/documents/{document_id}/process",
                    json={"task": "explain_code"},
                    headers={"X-User-API-Key": secret},
                )
        self.assertNotIn(secret, "\n".join(log_capture.output))
        self.assertEqual(
            self.client.get(
                f"/api/documents/{document_id}/tasks/{unexpected.json()['id']}"
            ).json()["error_code"],
            "PROCESSING_FAILED",
        )

        with patch.object(
            processing_module,
            "explain_code_blocks",
            new_callable=AsyncMock,
            return_value={f"code:{document_id}": "输出一行文本"},
        ):
            retried = self.client.post(
                f"/api/documents/{document_id}/process",
                json={"task": "explain_code"},
                headers={"X-User-API-Key": secret},
            )
        retried_status = self.client.get(
            f"/api/documents/{document_id}/tasks/{retried.json()['id']}"
        )
        self.assertEqual(retried_status.json()["status"], "completed")
        persisted_json = "\n".join(path.read_text(encoding="utf-8") for path in Path(self.temp_dir.name).rglob("*.json"))
        self.assertNotIn(secret, persisted_json)

    def test_interrupted_tasks_are_recoverable_and_path_ids_are_validated(self) -> None:
        document_id = "123e4567-e89b-42d3-a456-426614174000"
        task = create_processing_task(document_id, "translate")
        update_processing_task(task.id, status="processing", progress_percent=5)
        recover_interrupted_processing_tasks()
        self.assertEqual(get_document_task(document_id, task.id).status, "failed")
        with self.assertRaises(Exception):
            documents_module.get_parsed_result_paths("../../outside")

    def test_translation_preserves_code_formula_and_validates_ai_json(self) -> None:
        prepared = prepare_translation(
            "Compute `x = 1` and solve $x_i = 3$. The gradient is useful.",
            {"x", "x_i"},
        )
        translated = prepared.text.replace("Compute", "计算").replace("solve", "求解")
        restored = prepared.restore(translated)
        self.assertIn("`x = 1`", restored)
        self.assertIn("$x_i = 3$", restored)
        self.assertIn("梯度", restored)

        valid = TranslationOutput.model_validate_json(
            '{"items":[{"block_id":"block-1","translated_text":"翻译结果"}]}'
        )
        self.assertEqual(valid.items[0].block_id, "block-1")
        with self.assertRaises(ValidationError):
            TranslationOutput.model_validate_json(
                '{"items":[{"block_id":"block-1","translated_text":"ok","extra":true}]}'
            )

    def test_docling_docx_integration_extracts_table_and_code_blocks(self) -> None:
        source_path = settings.uploads_path / "layout-regression.docx"
        source_path.parent.mkdir(parents=True, exist_ok=True)
        fixture = Document()
        fixture.add_heading("Homework", 1)
        fixture.add_paragraph("A formula example: x_i = 3")
        fixture.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        fixture.add_paragraph("def square(x): return x ** 2", style="Code")
        table = fixture.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "method"
        table.cell(0, 1).text = "value"
        table.cell(1, 0).text = "square"
        table.cell(1, 1).text = "x ** 2"
        fixture.save(source_path)
        try:
            parsed = documents_module.parse_document(source_path, str(uuid4()))
        finally:
            source_path.unlink(missing_ok=True)
        self.assertEqual(parsed.status, "success")
        self.assertTrue(any(block.kind == "code" for block in parsed.parsed_blocks.blocks))
        self.assertTrue(any(block.kind == "table" for block in parsed.parsed_blocks.blocks))
        self.assertTrue(parsed.workbench_result.code_blocks)

    def test_docling_pdf_integration_reads_both_text_columns(self) -> None:
        source_path = settings.uploads_path / "two-columns.pdf"
        source_path.parent.mkdir(parents=True, exist_ok=True)
        source_path.write_bytes(make_two_column_pdf())
        try:
            parsed = documents_module.parse_document(source_path, str(uuid4()))
        finally:
            source_path.unlink(missing_ok=True)
        self.assertEqual(parsed.status, "success")
        for marker in (
            "LEFT-COLUMN-ONE",
            "LEFT-COLUMN-TWO",
            "RIGHT-COLUMN-ONE",
            "RIGHT-COLUMN-TWO",
        ):
            self.assertIn(marker, parsed.markdown)

    def test_docling_ocr_integration_reads_a_scanned_pdf(self) -> None:
        image = Image.new("RGB", (600, 800), "white")
        ImageDraw.Draw(image).text(
            (40, 80),
            "SCANNED PDF OCR TEST",
            fill="black",
            font=ImageFont.load_default(size=48),
        )
        jpeg = BytesIO()
        image.save(jpeg, format="JPEG", quality=95)
        source_path = settings.uploads_path / "scanned-page.pdf"
        source_path.parent.mkdir(parents=True, exist_ok=True)
        source_path.write_bytes(make_scanned_pdf(jpeg.getvalue(), 600, 800))
        try:
            parsed = documents_module.parse_document(source_path, str(uuid4()))
        finally:
            source_path.unlink(missing_ok=True)
        self.assertEqual(parsed.status, "success")
        self.assertIn("SCANNED", parsed.markdown.upper())

    @unittest.skipIf(os.name == "nt", "POSIX mode bits are not enforced on Windows")
    def test_local_documents_and_task_files_have_private_permissions(self) -> None:
        with patch.object(
            documents_module,
            "parse_document",
            side_effect=lambda _path, document_id: self.parsed_document(document_id),
        ):
            response = self.upload()
        self.assertEqual(response.status_code, 202)
        uploads = settings.uploads_path
        parsed = settings.parsed_path
        self.assertEqual(uploads.stat().st_mode & 0o777, 0o700)
        self.assertEqual(parsed.stat().st_mode & 0o777, 0o700)
        self.assertTrue(all((path.stat().st_mode & 0o777) == 0o600 for path in uploads.iterdir()))


if __name__ == "__main__":
    unittest.main()
